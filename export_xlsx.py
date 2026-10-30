"""Build an .xlsx report from a finished load-test run (bench.Benchmark.status())."""

import io
from datetime import datetime, timezone

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BACKENDS = (("opensearch", "OpenSearch"), ("s3vectors", "S3 Vectors"))
STATS = ("count", "avg", "min", "p50", "p90", "p95", "p99", "max")

HEADER_FONT = Font(bold=True, color="FFFFFF")
HEADER_FILL = PatternFill("solid", fgColor="305496")
TITLE_FONT = Font(bold=True, size=14)
BOLD = Font(bold=True)


def _ts(epoch):
    if not epoch:
        return ""
    return datetime.fromtimestamp(epoch, tz=timezone.utc).strftime("%Y-%m-%d %H:%M:%S UTC")


def _header(ws, row, values):
    for col, value in enumerate(values, start=1):
        cell = ws.cell(row=row, column=col, value=value)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def _autosize(ws, minimum=10, maximum=60):
    widths = {}
    for row in ws.iter_rows():
        for cell in row:
            if cell.value is not None:
                widths[cell.column] = max(widths.get(cell.column, 0), len(str(cell.value)))
    for col, width in widths.items():
        ws.column_dimensions[get_column_letter(col)].width = max(minimum, min(maximum, width + 2))


def _fusion(settings):
    pipeline = settings.get("pipeline")
    weight = settings.get("bm25_weight")
    if weight is None:
        return pipeline
    return f"{pipeline} (BM25 {weight:g} / kNN {round(1 - weight, 4):g})"


def build_workbook(status, queries):
    """Return the report as bytes."""
    s = status.get("settings", {})
    results = status.get("results", {})
    wb = Workbook()

    # -- Summary -------------------------------------------------------------
    ws = wb.active
    ws.title = "Summary"
    ws["A1"] = "Load test: S3 Vectors vs OpenSearch hybrid search"
    ws["A1"].font = TITLE_FONT
    info = [
        ("Run ID", status.get("id")),
        ("State", status.get("state")),
        ("Started", _ts(status.get("started"))),
        ("Finished", _ts(status.get("finished"))),
        ("Concurrency (threads)", s.get("concurrency")),
        ("Duration per backend (s)", s.get("duration_s")),
        ("Serial rounds", s.get("serial_rounds")),
        ("Queries", s.get("queries")),
        ("topK", s.get("k")),
        ("OpenSearch fusion", _fusion(s)),
        ("Warmup", "yes" if s.get("warmup") else "no"),
        ("Latency", "Client-side wall-clock time per search call (ms); query embedding excluded"),
    ]
    for i, (key, value) in enumerate(info, start=3):
        ws.cell(row=i, column=1, value=key).font = BOLD
        ws.cell(row=i, column=2, value=value)

    row = len(info) + 4
    _header(ws, row, ["Backend", "Serial avg latency (ms)", "Concurrent QPS",
                      "p50 (ms)", "p90 (ms)", "p95 (ms)", "p99 (ms)",
                      "OK requests", "Failed requests", "Throttled", "Error rate (%)"])
    for key, label in BACKENDS:
        r = results.get(key, {})
        serial = (r.get("serial") or {}).get("latency_ms", {})
        conc = r.get("concurrent") or {}
        lat = conc.get("latency_ms", {})
        row += 1
        ws.append([label, serial.get("avg"), conc.get("qps"),
                   lat.get("p50"), lat.get("p90"), lat.get("p95"), lat.get("p99"),
                   conc.get("requests_ok"), conc.get("requests_failed"),
                   conc.get("throttled"), conc.get("error_rate_pct")])
    _autosize(ws)
    ws.column_dimensions["A"].width = 26
    ws.column_dimensions["B"].width = max(ws.column_dimensions["B"].width, 22)

    # -- Latency details -------------------------------------------------------
    ws = wb.create_sheet("Latency details")
    _header(ws, 1, ["Backend", "Phase", "Metric"] + [f"{name} (ms)" if name != "count" else "count"
                                                    for name in STATS])
    for key, label in BACKENDS:
        r = results.get(key, {})
        for phase in ("serial", "concurrent"):
            data = r.get(phase) or {}
            for metric, field in (("client latency", "latency_ms"),
                                  ("server took", "server_took_ms")):
                stats = data.get(field)
                if not stats:
                    continue
                ws.append([label, phase, metric] + [stats.get(name) for name in STATS])
    _autosize(ws)

    # -- Errors ----------------------------------------------------------------
    ws = wb.create_sheet("Errors")
    _header(ws, 1, ["Backend", "Phase", "Error", "Count"])
    any_error = False
    for key, label in BACKENDS:
        r = results.get(key, {})
        for phase in ("serial", "concurrent"):
            for err, count in ((r.get(phase) or {}).get("errors") or {}).items():
                ws.append([label, phase, err, count])
                any_error = True
    if not any_error:
        ws.append(["(no errors)"])
    _autosize(ws)

    # -- Queries ---------------------------------------------------------------
    ws = wb.create_sheet("Queries")
    _header(ws, 1, ["#", "Query"])
    for i, q in enumerate(queries, start=1):
        ws.append([i, q])
    _autosize(ws, maximum=90)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
